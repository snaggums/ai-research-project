from pathlib import Path

import pytest
from app.services.parsing_service import (
    DOCX_TRANSCRIPT_PARSER,
    GENERIC_PARSER,
    extract_transcript,
)
from docx import Document


def _add_system_event(document: Document, text: str) -> None:
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 1).paragraphs[0].add_run(text)


def _add_turn(document: Document, speaker: str, timestamp: str, text: str) -> None:
    table = document.add_table(rows=1, cols=2)
    content = table.cell(0, 1)
    heading = content.paragraphs[0]
    speaker_run = heading.add_run(speaker)
    speaker_run.bold = True
    heading.add_run(f"  {timestamp}")
    content.add_paragraph(text)


def _templated_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("Session S003 - Meeting Recording")
    document.add_paragraph("July 10, 2026")
    document.add_paragraph("44m 0s")
    _add_system_event(document, "Maya Chen started transcription")
    _add_turn(
        document,
        "Maya Chen",
        "0:00",
        "Please send a secure message about your parent's medication.",
    )
    _add_turn(
        document,
        "Tanya",
        "3:06",
        "I am looking for Messages. I want to confirm I am in my mother's account.",
    )
    _add_system_event(document, "Maya Chen stopped transcription")
    document.save(path)


def test_extract_templated_docx_preserves_turn_boundaries_and_offsets(tmp_path: Path) -> None:
    path = tmp_path / "S003_Simplified_Transcript.docx"
    _templated_docx(path)

    parsed = extract_transcript(str(path))

    assert parsed.parser_name == DOCX_TRANSCRIPT_PARSER
    assert parsed.metadata == {
        "session_title": "Session S003 - Meeting Recording",
        "session_date": "July 10, 2026",
        "duration": "44m 0s",
        "source_turn_count": 2,
        "excluded_system_event_count": 2,
    }
    assert [(block.speaker, block.location, block.start_ms) for block in parsed.blocks] == [
        ("Maya Chen", "0:00", 0),
        ("Tanya", "3:06", 186000),
    ]
    assert all(
        parsed.content[block.start_char : block.end_char] == block.text
        for block in parsed.blocks
    )
    assert "started transcription" not in parsed.content
    assert "stopped transcription" not in parsed.content


def test_extract_generic_docx_includes_text_inside_tables(tmp_path: Path) -> None:
    path = tmp_path / "generic-notes.docx"
    document = Document()
    document.add_paragraph("Research notes")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).paragraphs[0].add_run("A finding stored in a table.")
    document.save(path)

    parsed = extract_transcript(str(path))

    assert parsed.parser_name == GENERIC_PARSER
    assert parsed.blocks == ()
    assert parsed.content == "Research notes\n\nA finding stored in a table."


def test_recognized_template_without_valid_turns_fails_actionably(tmp_path: Path) -> None:
    path = tmp_path / "malformed-transcript.docx"
    document = Document()
    document.add_paragraph("Session S009 - Meeting Recording")
    document.add_paragraph("July 20, 2026")
    document.add_paragraph("30m 0s")
    _add_system_event(document, "Maya Chen started transcription")
    document.save(path)

    with pytest.raises(ValueError, match="could not find any speaker turns"):
        extract_transcript(str(path))
