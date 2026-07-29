from io import BytesIO
from pathlib import Path
import re

import pytest
from app.models.document import Document as DocumentModel
from app.models.transcript_block import TranscriptBlockRecord
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session


pytestmark = pytest.mark.integration

PRIYA_NAIR_TRANSCRIPT = (
    Path(__file__).resolve().parents[2]
    / "sample-data"
    / "Sky_AIR_15_Synthetic_Fraud_Product_Interviews"
    / "product-1-medicare-fraud-documenter"
    / "P1-S01_Medicare_Fraud_Documenter_Priya_Nair.docx"
)


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Session S003 - Meeting Recording")
    document.add_paragraph("July 10, 2026")
    document.add_paragraph("44m 0s")

    start = document.add_table(rows=1, cols=2)
    start.cell(0, 1).paragraphs[0].add_run("Maya Chen started transcription")

    for speaker, timestamp, text in (
        (
            "Maya Chen",
            "0:00",
            "Please send a secure message about your parent's medication.",
        ),
        (
            "Tanya",
            "3:06",
            "I am looking for Messages. I want to confirm I am in my mother's account.",
        ),
    ):
        table = document.add_table(rows=1, cols=2)
        content = table.cell(0, 1)
        speaker_run = content.paragraphs[0].add_run(speaker)
        speaker_run.bold = True
        content.paragraphs[0].add_run(f"  {timestamp}")
        content.add_paragraph(text)

    stop = document.add_table(rows=1, cols=2)
    stop.cell(0, 1).paragraphs[0].add_run("Maya Chen stopped transcription")

    value = BytesIO()
    document.save(value)
    return value.getvalue()


def _workspace(client: TestClient) -> tuple[dict, dict]:
    project_response = client.post(
        "/api/projects",
        json={"name": "Structured transcript project", "description": "DOCX ingestion"},
    )
    assert project_response.status_code == 201
    project = project_response.json()
    session_response = client.post(
        f"/api/projects/{project['id']}/sessions",
        json={
            "title": "Structured transcript session",
            "type": "interview",
            "participant_ids": [],
            "related_record_ids": ["record-1"],
        },
    )
    assert session_response.status_code == 201
    return project, session_response.json()


def test_templated_docx_flows_to_transcript_coding_and_search(
    client: TestClient,
    db_session: Session,
) -> None:
    project, research_session = _workspace(client)
    root = f"/api/projects/{project['id']}/sessions/{research_session['id']}"
    upload = client.post(
        f"{root}/documents",
        files={
            "file": (
                "S003_Simplified_Transcript.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == 201
    document_id = upload.json()["id"]

    listed = client.get(f"{root}/documents")
    assert listed.status_code == 200
    document = next(value for value in listed.json() if value["id"] == document_id)
    assert document["status"] == "complete"
    assert [(block["speaker"], block["location"]) for block in document["blocks"]] == [
        ("Maya Chen", "0:00"),
        ("Tanya", "3:06"),
    ]

    stored_document = db_session.get(DocumentModel, document_id)
    assert stored_document is not None
    assert stored_document.parser_name == "docx-templated-transcript"
    stored_blocks = list(
        db_session.scalars(
            select(TranscriptBlockRecord)
            .where(TranscriptBlockRecord.document_id == document_id)
            .order_by(TranscriptBlockRecord.block_index)
        ).all()
    )
    assert [value.id for value in stored_blocks] == [value["id"] for value in document["blocks"]]

    coding = client.get(f"{root}/coding")
    assert coding.status_code == 200
    workspace = coding.json()
    assert [(block["speaker"], block["location"]) for block in workspace["transcript"]["blocks"]] == [
        ("Maya Chen", "0:00"),
        ("Tanya", "3:06"),
    ]
    first_block = workspace["transcript"]["blocks"][0]
    assert first_block["text"] == document["blocks"][0]["text"]

    retry_before_coding = client.post(f"{root}/documents/{document_id}/process")
    assert retry_before_coding.status_code == 200
    refreshed_coding = client.get(f"{root}/coding")
    assert refreshed_coding.status_code == 200
    assert [block["id"] for block in refreshed_coding.json()["transcript"]["blocks"]] == [
        block["id"] for block in workspace["transcript"]["blocks"]
    ]
    workspace = refreshed_coding.json()
    first_block = workspace["transcript"]["blocks"][0]

    search = client.get(
        f"{root}/documents/{document_id}/search",
        params={"q": "secure message"},
    )
    assert search.status_code == 200
    result = search.json()["results"][0]
    assert result["speaker"] == "Maya Chen"
    assert result["location"] == "0:00"

    context = client.get(f"{root}/documents/{document_id}/context/{result['id']}")
    assert context.status_code == 200
    assert context.json()["focused_passage_id"] == first_block["id"]

    highlight = client.post(
        f"{root}/highlights",
        headers={"Idempotency-Key": "structured-docx-highlight"},
        json={
            "anchor": {
                "chunk_id": first_block["chunk_id"],
                "block_id": first_block["id"],
                "start_char": first_block["start_char"],
                "end_char": first_block["end_char"],
                "excerpt_snapshot": first_block["text"],
                "speaker": first_block["speaker"],
                "location": first_block["location"],
                "start_ms": 0,
                "content_checksum": workspace["transcript"]["content_checksum"],
            },
            "code_ids": [],
            "new_code": None,
        },
    )
    assert highlight.status_code == 201, highlight.text
    assert highlight.json()["anchor"]["block_id"] == first_block["id"]

    retry = client.post(f"{root}/documents/{document_id}/process")
    assert retry.status_code == 409
    assert retry.json()["code"] == "transcript_has_coding_dependencies"


def test_transcript_search_returns_only_exact_billing_matches(client: TestClient) -> None:
    project, research_session = _workspace(client)
    root = f"/api/projects/{project['id']}/sessions/{research_session['id']}"
    upload = client.post(
        f"{root}/documents",
        files={
            "file": (
                PRIYA_NAIR_TRANSCRIPT.name,
                PRIYA_NAIR_TRANSCRIPT.read_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert upload.status_code == 201
    document_id = upload.json()["id"]

    response = client.get(
        f"{root}/documents/{document_id}/search",
        params={"q": "BILLING"},
    )

    assert response.status_code == 200
    results = response.json()["results"]
    assert len(results) == 3
    assert [result["block_index"] for result in results] == sorted(
        result["block_index"] for result in results
    )
    assert all(
        re.search(r"(?<!\w)billing(?!\w)", f"{result['speaker']} {result['excerpt']}", re.IGNORECASE)
        for result in results
    )

    no_match = client.get(
        f"{root}/documents/{document_id}/search",
        params={"q": "rebilling"},
    )
    assert no_match.status_code == 200
    assert no_match.json()["results"] == []
