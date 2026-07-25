from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator

from docx import Document as DocxDocument
from docx.document import Document as DocxDocumentType
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".txt", ".md", ".docx", ".pdf"}
DOCX_TRANSCRIPT_PARSER = "docx-templated-transcript"
DOCX_TRANSCRIPT_PARSER_VERSION = "1"
GENERIC_PARSER = "generic-text"
GENERIC_PARSER_VERSION = "1"

_TIMESTAMP_PATTERN = re.compile(r"(?P<timestamp>\d{1,2}:\d{2}(?::\d{2})?)\s*$")


@dataclass(frozen=True)
class ParsedTranscriptBlock:
    block_index: int
    speaker: str
    location: str
    text: str
    start_char: int
    end_char: int
    start_ms: int | None = None
    end_ms: int | None = None
    source_metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class ParsedTranscript:
    content: str
    blocks: tuple[ParsedTranscriptBlock, ...]
    parser_name: str
    parser_version: str
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class _DocxTurn:
    speaker: str
    timestamp: str
    text: str
    table_index: int


def extract_transcript(file_path: str) -> ParsedTranscript:
    path = Path(file_path)
    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type. Upload one of: {supported}")

    if extension in {".txt", ".md"}:
        content = path.read_text(encoding="utf-8", errors="replace").strip()
        return ParsedTranscript(content, (), GENERIC_PARSER, GENERIC_PARSER_VERSION)

    if extension == ".docx":
        return _extract_docx_transcript(path)

    reader = PdfReader(path)
    pages = [page.extract_text() or "" for page in reader.pages]
    content = "\n\n".join(page.strip() for page in pages if page.strip()).strip()
    return ParsedTranscript(
        content,
        (),
        GENERIC_PARSER,
        GENERIC_PARSER_VERSION,
        {"page_count": len(reader.pages)},
    )


def extract_text(file_path: str) -> str:
    return extract_transcript(file_path).content


def _extract_docx_transcript(path: Path) -> ParsedTranscript:
    document = DocxDocument(path)
    header_paragraphs = [_normalize_text(paragraph.text) for paragraph in document.paragraphs]
    header_paragraphs = [value for value in header_paragraphs if value]
    generic_parts: list[str] = []
    turns: list[_DocxTurn] = []
    system_event_count = 0
    table_index = 0

    for item in _iter_docx_body_items(document):
        if isinstance(item, Paragraph):
            text = _normalize_text(item.text)
            if text:
                generic_parts.append(text)
            continue

        turn = _parse_templated_turn(item, table_index)
        if turn is not None:
            turns.append(turn)
        else:
            table_text = _table_text(item)
            if table_text:
                generic_parts.append(table_text)
                if _is_transcription_event(table_text):
                    system_event_count += 1
        table_index += 1

    if turns:
        content, blocks = _normalize_turns(turns)
        metadata: dict[str, Any] = {
            "session_title": header_paragraphs[0] if len(header_paragraphs) > 0 else None,
            "session_date": header_paragraphs[1] if len(header_paragraphs) > 1 else None,
            "duration": header_paragraphs[2] if len(header_paragraphs) > 2 else None,
            "source_turn_count": len(turns),
            "excluded_system_event_count": system_event_count,
        }
        return ParsedTranscript(
            content=content,
            blocks=blocks,
            parser_name=DOCX_TRANSCRIPT_PARSER,
            parser_version=DOCX_TRANSCRIPT_PARSER_VERSION,
            metadata=metadata,
        )

    looks_like_templated_transcript = (
        bool(header_paragraphs)
        and header_paragraphs[0].casefold().endswith("meeting recording")
        and table_index > 0
    )
    if looks_like_templated_transcript:
        raise ValueError(
            "AIR recognized this meeting Transcript template but could not find any speaker turns with timestamps."
        )

    content = "\n\n".join(generic_parts).strip()
    return ParsedTranscript(
        content=content,
        blocks=(),
        parser_name=GENERIC_PARSER,
        parser_version=GENERIC_PARSER_VERSION,
        metadata={"body_item_count": len(generic_parts)},
    )


def _iter_docx_body_items(document: DocxDocumentType) -> Iterator[Paragraph | Table]:
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _parse_templated_turn(table: Table, table_index: int) -> _DocxTurn | None:
    if not table.rows or len(table.rows[0].cells) < 2:
        return None

    content_cell = table.rows[0].cells[-1]
    paragraphs = [paragraph for paragraph in content_cell.paragraphs if _normalize_text(paragraph.text)]
    if len(paragraphs) < 2:
        return None

    heading = _normalize_text(paragraphs[0].text)
    timestamp_match = _TIMESTAMP_PATTERN.search(heading)
    if timestamp_match is None:
        return None

    bold_speaker = "".join(
        run.text for run in paragraphs[0].runs if run.bold and run.text
    )
    speaker = _normalize_inline_text(bold_speaker)
    if not speaker:
        speaker = _normalize_inline_text(heading[: timestamp_match.start()])
    if not speaker:
        return None

    body = "\n".join(
        value
        for value in (_normalize_text(paragraph.text) for paragraph in paragraphs[1:])
        if value
    )
    if not body:
        return None

    return _DocxTurn(
        speaker=speaker,
        timestamp=timestamp_match.group("timestamp"),
        text=body,
        table_index=table_index,
    )


def _normalize_turns(
    turns: list[_DocxTurn],
) -> tuple[str, tuple[ParsedTranscriptBlock, ...]]:
    content_parts: list[str] = []
    blocks: list[ParsedTranscriptBlock] = []
    cursor = 0

    for block_index, turn in enumerate(turns):
        if content_parts:
            content_parts.append("\n\n")
            cursor += 2

        heading = f"{turn.speaker} [{turn.timestamp}]\n"
        content_parts.append(heading)
        cursor += len(heading)
        start_char = cursor

        content_parts.append(turn.text)
        cursor += len(turn.text)
        blocks.append(
            ParsedTranscriptBlock(
                block_index=block_index,
                speaker=turn.speaker,
                location=turn.timestamp,
                text=turn.text,
                start_char=start_char,
                end_char=cursor,
                start_ms=_timestamp_to_ms(turn.timestamp),
                source_metadata={"source_table_index": turn.table_index},
            )
        )

    return "".join(content_parts), tuple(blocks)


def _table_text(table: Table) -> str:
    values: list[str] = []
    seen_cells: set[int] = set()
    for row in table.rows:
        for cell in row.cells:
            cell_identity = id(cell._tc)
            if cell_identity in seen_cells:
                continue
            seen_cells.add(cell_identity)
            for paragraph in cell.paragraphs:
                text = _normalize_text(paragraph.text)
                if text:
                    values.append(text)
    return "\n".join(values)


def _is_transcription_event(value: str) -> bool:
    normalized = value.casefold()
    return normalized.endswith("started transcription") or normalized.endswith("stopped transcription")


def _timestamp_to_ms(value: str) -> int:
    parts = [int(part) for part in value.split(":")]
    if len(parts) == 2:
        minutes, seconds = parts
        return (minutes * 60 + seconds) * 1000
    hours, minutes, seconds = parts
    return (hours * 3600 + minutes * 60 + seconds) * 1000


def _normalize_text(value: str) -> str:
    normalized = unicodedata.normalize("NFC", value).replace("\r\n", "\n").replace("\r", "\n")
    lines = [_normalize_inline_text(line) for line in normalized.split("\n")]
    return "\n".join(line for line in lines if line).strip()


def _normalize_inline_text(value: str) -> str:
    return " ".join(unicodedata.normalize("NFC", value).split())
